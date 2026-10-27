# -*- coding: utf-8 -*-
"Telegram-бот @test2_for_work_bot"
import asyncio
import csv
import json
import logging
import math
import os
import re
from datetime import datetime
from io import BytesIO
from fnmatch import fnmatch
from pathlib import Path

from aiogram import Bot, Dispatcher, F, Router
from aiogram.exceptions import TelegramBadRequest
from aiogram.filters import Command, CommandStart
from aiogram.types import (BufferedInputFile, BotCommand, BotCommandScopeChat, BotCommandScopeDefault, CallbackQuery,
                           FSInputFile, InlineKeyboardButton, InlineKeyboardMarkup, KeyboardButton,
                           MenuButtonCommands, Message, ReplyKeyboardMarkup, ReplyKeyboardRemove)
from dotenv import load_dotenv

from scenario import BACK, GROUPS, STEPS

BASE = Path(__file__).parent
DATA = BASE / "data"
DATA.mkdir(exist_ok=True)
OVERRIDES_FILE, SETTINGS_FILE = DATA / "overrides.json", DATA / "settings.json"
LEADS_FILE, PRICES_FILE = DATA / "leads.csv", BASE / "prices.json"

log = logging.getLogger("bot")
router = Router()
runtime_dispatcher = None
restart_requested = False


# ───────────────────────── Хранилище ─────────────────────────
def load_json(path, default):
    try:
        return json.loads(path.read_text("utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return default


def save_json(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), "utf-8")


overrides: dict = load_json(OVERRIDES_FILE, {})   # тексты, изменённые через /admin
settings: dict = load_json(SETTINGS_FILE, {})     # chat_id менеджера и т.п.


def step_text(step):
    return overrides.get(step["id"], step["text"])


def privacy_confirmed(uid):
    accepted = settings.get("privacy_accepted", [])
    return isinstance(accepted, list) and str(uid) in {str(value) for value in accepted}


def is_admin(uid):
    ids = {x for x in re.split(r"[,\s]+", os.getenv("ADMIN_IDS", "")) if x}
    stored = settings.get("admin_ids", [])
    if isinstance(stored, list):
        ids.update(str(x) for x in stored)
    return str(uid) in ids


def manager_chat():
    return settings.get("manager_chat") or os.getenv("MANAGER_CHAT_ID") or None


def is_manager(uid):
    """Менеджер: тот, чей личный чат указан для заявок, либо ID из MANAGER_IDS в .env."""
    ids = {x for x in re.split(r"[,\s]+", os.getenv("MANAGER_IDS", "")) if x}
    stored = settings.get("manager_ids", [])
    if isinstance(stored, list):
        ids.update(str(x) for x in stored)
    mc = manager_chat()
    return str(uid) in ids or (mc is not None and str(mc) == str(uid))


# ───────────────────────── Меню бота (кнопка «Меню» рядом с полем ввода) ─────────────────────────
def menu_commands(uid, in_branch):
    if not privacy_confirmed(uid):
        return []
    cmds = [BotCommand(command="start", description="Вернуться к началу" if in_branch else "Начать"),
            BotCommand(command="contacts", description="Контакты")]
    if is_manager(uid) or is_admin(uid):
        cmds.append(BotCommand(command="leads", description="Последние заявки"))
    if is_admin(uid):
        cmds.append(BotCommand(command="admin", description="Админ-панель"))
    return cmds


async def update_menu(bot: Bot, chat_id, s, force=False):
    """«Начать» → «Вернуться к началу», как только клиент ушёл в какую-либо ветку."""
    uid, st = s["user"]["id"], s["stack"]
    in_branch = bool(st) and st[-1] not in ("privacy", "main")
    key = (privacy_confirmed(uid), in_branch, is_manager(uid), is_admin(uid))
    if not force and s.get("menu_key") == key:
        return
    try:
        await bot.set_my_commands(menu_commands(uid, in_branch), scope=BotCommandScopeChat(chat_id=chat_id))
        s["menu_key"] = key
    except Exception:
        log.exception("Не удалось обновить меню команд")


async def clear_kb(m: Message):
    """Убираем старую нижнюю клавиатуру (если осталась с прошлой версии бота)."""
    carrier = await m.bot.send_message(m.chat.id, "\u2063", reply_markup=ReplyKeyboardRemove(),
                                       disable_notification=True)
    asyncio.create_task(delete_keyboard_carrier(m.bot, m.chat.id, carrier.message_id))


async def delete_keyboard_carrier(bot: Bot, chat_id, message_id):
    await asyncio.sleep(2)
    try:
        await bot.delete_message(chat_id, message_id)
    except TelegramBadRequest:
        log.debug("Не удалось удалить временное сообщение клавиатуры")


async def remove_menu_editor_keyboard(msg: Message, s):
    if not s.pop("editor_keyboard_message", None):
        return
    await clear_kb(msg)


# ───────────────────────── Заявки: хранилище ─────────────────────────
LEAD_COLS = ["ID", "Дата", "Имя", "Телефон", "Мессенджеры", "Время", "Telegram", "Ветка", "Оценка", "Ответы"]
LEADS_PAGE = 5


def write_leads(leads):
    with LEADS_FILE.open("w", newline="", encoding="utf-8-sig") as f:   # utf-8-sig — чтобы Excel не ломал кириллицу
        w = csv.writer(f, delimiter=";")
        w.writerow(LEAD_COLS)
        for l in leads:
            w.writerow([l[c] for c in LEAD_COLS])


def read_leads():
    try:
        with LEADS_FILE.open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f, delimiter=";"))
    except FileNotFoundError:
        return []
    if not rows:
        return []
    head, body = rows[0], rows[1:]
    old_format = not head or head[0] != "ID"
    if old_format:                                  # файл со старой версии бота: нумеруем по порядку
        body = [[str(i)] + r for i, r in enumerate(body, 1)]
    leads = [dict(zip(LEAD_COLS, (r + [""] * len(LEAD_COLS))[:len(LEAD_COLS)])) for r in body]
    if old_format:
        write_leads(leads)
        settings["last_lead_id"] = max([int(l["ID"]) for l in leads] + [settings.get("last_lead_id", 0)])
        save_json(SETTINGS_FILE, settings)
    return leads


def next_lead_id(leads):
    """ID растёт всегда: номер удалённой заявки не выдаётся повторно."""
    nid = max([int(l["ID"]) for l in leads if l["ID"].isdigit()] + [settings.get("last_lead_id", 0)]) + 1
    settings["last_lead_id"] = nid
    save_json(SETTINGS_FILE, settings)
    return nid


def lead_card(l):
    out = [f"🆔 Заявка №{l['ID']} · {l['Дата']}", "", f"👤 {l['Имя']}", f"📞 {l['Телефон']}",
           f"💬 Мессенджеры: {l['Мессенджеры'] or '—'}", f"🕐 Удобное время: {l['Время'] or '—'}",
           f"✈️ Telegram: {l['Telegram']}", f"📌 Ветка: {l['Ветка']}"]
    if l["Оценка"]:
        out.append(f"💰 Вилка: {l['Оценка']}")
    if l["Ответы"]:
        out += ["", "Ответы клиента:"] + l["Ответы"].split(" | ")
    return "\n".join(out)[:3900]


def build_xlsx(leads) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявки"
    ws.append(LEAD_COLS)
    for l in leads:
        ws.append([int(l["ID"]) if l["ID"].isdigit() else l["ID"]]
                  + [l[c] for c in LEAD_COLS[1:-1]] + [l["Ответы"].replace(" | ", "\n")])
    for i, w in enumerate([6, 17, 20, 18, 22, 18, 18, 28, 22, 70]):
        ws.column_dimensions[chr(65 + i)].width = w
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows():
        for c in row:
            c.alignment = wrap
            if isinstance(c.value, str):
                c.data_type = "s"                   # текст клиента не должен выполняться как формула
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ───────────────────────── Сессии ─────────────────────────
sessions: dict = {}


def sess(user):
    s = sessions.setdefault(user.id, {"stack": [], "ans": {}, "files": [], "ask": None,
                                      "submitted": False, "adm_edit": None})
    s["user"] = {"id": user.id, "username": user.username, "name": user.full_name}
    return s


def push(s, sid):
    st = s["stack"]
    if sid in st:                       # возврат к уже пройденному шагу (например, к списку разделов)
        del st[st.index(sid) + 1:]
    else:
        st.append(sid)


def client_name(s):
    a = s["ans"].get("f_name")
    return a["sel"][0] if a and a["sel"] else ""


def menu_choice_config(step, create=False):
    choices = settings.get("menu_choices", {})
    config = choices.get(step["id"], {})
    if create:
        choices = settings.setdefault("menu_choices", {})
        config = choices.setdefault(step["id"], {})
        config.setdefault("hidden", [])
        config.setdefault("deleted", [])
        config.setdefault("labels", {})
        config.setdefault("added", [])
    return config


def effective_options(step, include_hidden=False):
    config = menu_choice_config(step)
    options = []
    for i, original in enumerate(step["options"]):
        key = f"base_{i}"
        if key in config.get("deleted", []) or (key in config.get("hidden", []) and not include_hidden):
            continue
        option = dict(original)
        option["label"] = config.get("labels", {}).get(key, option["label"])
        option["_menu_key"] = key
        option["_hidden"] = key in config.get("hidden", [])
        options.append(option)
    for original in config.get("added", []):
        key = original["id"]
        if key in config.get("deleted", []) or (key in config.get("hidden", []) and not include_hidden):
            continue
        option = dict(original)
        option["label"] = config.get("labels", {}).get(key, option["label"])
        option["_menu_key"] = key
        option["_hidden"] = key in config.get("hidden", [])
        options.append(option)
    return options


def allows_added_choices(step):
    options = effective_options(step, include_hidden=True)
    if not step.get("record", True) or not any(option.get("rec", True) for option in options):
        return False
    has_sequential_question = isinstance(step.get("next"), str) and step["next"] in STEPS
    return has_sequential_question or all(option.get("rec", True) for option in options)


def keyboard(step, s):
    rows, a = [], s["ans"].get(step["id"], {})
    for i, opt in enumerate(effective_options(step)):
        label = opt["label"]
        if step["type"] == "multi":
            label = ("✅ " if opt["label"] in a.get("sel", []) else "⬜ ") + label
        rows.append([InlineKeyboardButton(text=label, callback_data=f"o:{opt['_menu_key']}")])
    if step["type"] == "multi":
        rows.append([InlineKeyboardButton(text=step.get("done_label", "➡️ Далее"), callback_data="done")])
    if step.get("back", True) and len(s["stack"]) > 1:
        rows.append([InlineKeyboardButton(text=BACK, callback_data="back")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


async def render(msg: Message, s, text, markup):
    """Единственное меню: редактируем сохранённое сообщение, новое шлём только если старого нет."""
    mid = s.get("menu")
    if mid:
        try:
            await msg.bot.edit_message_text(text, chat_id=msg.chat.id, message_id=mid, reply_markup=markup)
            return
        except TelegramBadRequest as e:
            if "not modified" in str(e):
                return
    sent = await msg.bot.send_message(msg.chat.id, text, reply_markup=markup)
    s["menu"] = sent.message_id
    s["menu_chat_id"] = msg.chat.id
    track_chat_message(s, sent.message_id, msg.chat.id)


def track_chat_message(s, message_id, chat_id):
    chat_key = str(chat_id)
    tracked = s.setdefault("chat_messages", {}).setdefault(chat_key, [])
    if message_id not in tracked:
        tracked.append(message_id)
    cleanup = settings.setdefault("chat_cleanup", {})
    ids = cleanup.setdefault(chat_key, [])
    if message_id not in ids:
        ids.append(message_id)
        cleanup[chat_key] = ids[-100:]
        save_json(SETTINGS_FILE, settings)


async def clear_chat(msg: Message, s, preserve_current=False):
    if not preserve_current:
        await try_delete(msg)
    s["price_edit"] = None
    cleanup = settings.get("chat_cleanup", {})
    chat_key = str(msg.chat.id)
    persisted_ids = cleanup.get(chat_key, []) if isinstance(cleanup, dict) else []
    session_ids = s.get("chat_messages", {}).get(chat_key, [])
    message_ids = list(dict.fromkeys(
        session_ids + persisted_ids
        + ([s.get("menu")] if s.get("menu_chat_id") == msg.chat.id else [])
        + ([s.get("menu_hint")] if s.get("menu_hint_chat_id") == msg.chat.id else [])
    ))
    for message_id in message_ids:
        if message_id is None:
            continue
        try:
            await msg.bot.delete_message(msg.chat.id, message_id)
        except TelegramBadRequest:
            pass
    s.get("chat_messages", {}).pop(chat_key, None)
    if s.get("menu_chat_id") == msg.chat.id:
        s.pop("menu", None)
        s.pop("menu_chat_id", None)
    if s.get("menu_hint_chat_id") == msg.chat.id:
        s.pop("menu_hint", None)
        s.pop("menu_hint_chat_id", None)
    if isinstance(cleanup, dict):
        cleanup.pop(chat_key, None)
        save_json(SETTINGS_FILE, settings)
    if s.get("editor_keyboard_message"):
        await remove_menu_editor_keyboard(msg, s)


async def show_privacy(msg: Message, s):
    step = STEPS["privacy"]
    text = step_text(step)
    markup = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ Согласен, продолжить", callback_data="privacy:accept")]
    ])
    await render(msg, s, text, markup)
    await update_menu(msg.bot, msg.chat.id, s, force=True)


async def require_privacy(msg: Message, s):
    if privacy_confirmed(msg.from_user.id):
        return True
    await clear_chat(msg, s)
    await show_privacy(msg, s)
    return False


async def require_privacy_callback(cb: CallbackQuery, s):
    if privacy_confirmed(cb.from_user.id):
        return True
    await clear_chat(cb.message, s)
    await show_privacy(cb.message, s)
    await cb.answer("Сначала подтвердите политику конфиденциальности.", show_alert=True)
    return False


async def clear_menu_hint(msg: Message, s):
    if s.get("menu_hint_chat_id") != msg.chat.id:
        return
    hint_id = s.pop("menu_hint", None)
    s.pop("menu_hint_chat_id", None)
    if hint_id:
        try:
            await msg.bot.delete_message(msg.chat.id, hint_id)
        except TelegramBadRequest:
            pass


async def send_menu_hint(msg: Message, s, text, markup=None):
    await clear_menu_hint(msg, s)
    hint = await msg.bot.send_message(msg.chat.id, text, reply_markup=markup)
    s["menu_hint"] = hint.message_id
    s["menu_hint_chat_id"] = msg.chat.id
    track_chat_message(s, hint.message_id, msg.chat.id)


async def show(msg: Message, s, sid, notice=""):
    step = STEPS[sid]
    text = step_text(step).replace("{name}", client_name(s))
    if notice:
        text = f"{notice}\n\n{text}"
    s["ld_ask"] = False
    await render(msg, s, text, keyboard(step, s))
    await update_menu(msg.bot, msg.chat.id, s)


async def prompt(msg: Message, s, text):
    """Вопрос «Напишите…» показываем в том же меню, с кнопкой отмены."""
    await render(msg, s, text, InlineKeyboardMarkup(inline_keyboard=[btn(BACK, "cancelask")]))


async def try_delete(m: Message):
    try:
        await m.delete()
    except TelegramBadRequest:
        pass


TEMP_TTL = 15                           # через сколько секунд исчезают служебные сообщения
_bg: set = set()


async def _delete_later(*msgs: Message, delay=TEMP_TTL):
    await asyncio.sleep(delay)
    for x in msgs:
        await try_delete(x)


async def answer_temp(m: Message, text, delay=TEMP_TTL):
    """Служебное сообщение: показываем, а через `delay` секунд удаляем вместе с командой клиента."""
    sent = await m.bot.send_message(m.chat.id, text)
    track_chat_message(sess(m.from_user), sent.message_id, m.chat.id)
    task = asyncio.create_task(_delete_later(m, sent, delay=delay))
    _bg.add(task)
    task.add_done_callback(_bg.discard)
    return sent


async def go(msg: Message, s, sid, notice=""):
    if sid == "main":                   # главное меню = начать заявку заново
        s.update(stack=["main"], ans={}, files=[], ask=None, submitted=False)
    else:
        push(s, sid)
    if sid == "f_done" and not s["submitted"]:
        s["submitted"] = True
        await submit_lead(msg.bot, s)
    await show(msg, s, sid, notice)


def next_of(step, opt, s):
    n = (opt or {}).get("next") or step.get("next")
    return n(s) if callable(n) else n


# ───────────────────────── Заявка: сводка, оценка, отправка ─────────────────────────
FLOW_NAMES = {"p": "Новый сайт — товары", "s": "Новый сайт — услуги", "m": "Новый сайт — товары и услуги",
              "r": "Доработка сайта", "pr": "Продвижение сайта", "n": "Новый сайт (своё описание)",
              "ex": "Запрос примеров работ"}


def detect_flow(ans):
    prefixes = {k.split("_", 1)[0] for k in ans}
    for p in ("p", "s", "m", "r", "pr", "n", "ex"):
        if p in prefixes:
            return p
    return ""


def summary_lines(ans):
    lines = []
    for key, a in ans.items():
        if key.startswith("f_"):
            continue
        extra = a.get("extra", {})
        parts = [f"{x} ({extra[x]})" if x in extra else x for x in a["sel"]]
        parts += [f"{k}: {v}" for k, v in extra.items() if k not in a["sel"]]
        if parts:
            lines.append(f"• {a['title']}: {', '.join(parts)}")
    return lines


def estimate(ans):
    cfg = load_json(PRICES_FILE, {})
    code = detect_flow(ans)
    if code not in cfg.get("base", {}):
        return None
    total = cfg["base"][code]
    for key, a in ans.items():
        suffix = key.split("_", 1)[1]
        for r in cfg.get("rules", []):
            if not fnmatch(suffix, r["step"]):
                continue
            if "per_item" in r:
                total += r["per_item"] * len([x for x in a["sel"] if x.strip() != "Нет"])
            elif any(r["answer"] in x for x in a["sel"]):
                total += r["add"]
    lo, hi = cfg.get("range", [0.85, 1.3])
    rnd = lambda v: int(round(v / 5000) * 5000)
    fmt = lambda v: f"{v:,}".replace(",", " ")
    return f"{fmt(rnd(total * lo))} – {fmt(rnd(total * hi))} {cfg.get('currency', '₽')}"


def chunks(text, size=3900):
    buf = ""
    for line in text.split("\n"):
        if len(buf) + len(line) + 1 > size:
            yield buf
            buf = ""
        buf += line + "\n"
    if buf.strip():
        yield buf


async def submit_lead(bot: Bot, s):
    ans = s["ans"]
    first = lambda k: ans[k]["sel"][0] if k in ans and ans[k]["sel"] else ""
    name, phone, when = first("f_name"), first("f_phone"), first("f_time")
    messengers = ", ".join(ans.get("f_msg", {}).get("sel", []))
    flow = FLOW_NAMES.get(detect_flow(ans), "Обращение без выбранной ветки")
    est = estimate(ans)
    u = s["user"]
    tg = f"@{u['username']}" if u.get("username") else f"id {u['id']}"
    lines = summary_lines(ans)

    leads = read_leads()
    lead_id = next_lead_id(leads)
    card = (f"🔔 Новая заявка №{lead_id}\n\n👤 {name}\n📞 {phone}\n💬 Мессенджеры: {messengers or '—'}\n"
            f"🕐 Удобное время: {when or '—'}\n✈️ Telegram: {tg}\n📌 Ветка: {flow}\n")
    if est:
        card += f"💰 Предварительная вилка: {est} (ориентир, не оферта)\n"
    card += "\nОтветы клиента:\n" + "\n".join(lines)

    leads.append(dict(zip(LEAD_COLS, [str(lead_id), datetime.now().strftime("%Y-%m-%d %H:%M"), name, phone,
                                      messengers, when, tg, flow, est or "", " | ".join(lines)])))
    write_leads(leads)

    chat = manager_chat()
    if not chat:
        log.warning("Менеджер не задан: отправьте боту /setmanager. Заявка сохранена в data/leads.csv")
        return
    try:
        for part in chunks(card):
            await bot.send_message(chat, part)
        for kind, fid in s["files"]:
            await (bot.send_document if kind == "document" else bot.send_photo)(chat, fid)
    except Exception:
        log.exception("Не удалось отправить заявку менеджеру")


# ───────────────────────── Команды клиента ─────────────────────────
@router.message(CommandStart())
async def cmd_start(m: Message):
    s = sess(m.from_user)
    await clear_chat(m, s, preserve_current=True)
    if not privacy_confirmed(m.from_user.id):
        s.update(stack=[], ans={}, files=[], ask=None, submitted=False)
        await show_privacy(m, s)
        await try_delete(m)
        return
    s.update(stack=[], ans={}, files=[], ask=None, submitted=False)
    await go(m, s, "main")
    await try_delete(m)


@router.message(Command("menu"))
async def cmd_menu(m: Message):
    s = sess(m.from_user)
    if not await require_privacy(m, s):
        return
    await clear_chat(m, s)
    s["role_add"] = None
    s["menu_edit"] = None
    await go(m, s, "main")


@router.message(Command("myid"))
async def cmd_myid(m: Message):
    s = sess(m.from_user)
    if not await require_privacy(m, s):
        return
    await clear_chat(m, s)
    await answer_temp(m, f"Ваш Telegram ID: {m.from_user.id}\nID этого чата: {m.chat.id}", delay=60)


@router.callback_query(F.data == "privacy:accept")
async def cb_privacy_accept(cb: CallbackQuery):
    s = sess(cb.from_user)
    if not privacy_confirmed(cb.from_user.id):
        accepted = settings.setdefault("privacy_accepted", [])
        if not isinstance(accepted, list):
            accepted = settings["privacy_accepted"] = []
        accepted.append(str(cb.from_user.id))
        save_json(SETTINGS_FILE, settings)
    s["menu"] = cb.message.message_id
    s.update(stack=[], ans={}, files=[], ask=None, submitted=False)
    await go(cb.message, s, "main")
    await cb.answer("Спасибо за подтверждение.")


# ───────────────────────── Админ-панель ─────────────────────────
def admin_home():
    return ("🔧 Панель администратора", InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🧭 Редактировать меню и пункты", callback_data="adm:menu")],
        [InlineKeyboardButton(text="💰 Редактировать цены", callback_data="adm:prices")],
        [InlineKeyboardButton(text="↩️ Сбросить все меню и выборы", callback_data="adm:resetall")],
        [InlineKeyboardButton(text="📋 Заявки", callback_data="ld:l:0")],
        [InlineKeyboardButton(text="📬 Куда приходят заявки", callback_data="adm:mgr")],
        [InlineKeyboardButton(text="➕ Добавить администратора", callback_data="adm:addadmin")],
        [InlineKeyboardButton(text="➕ Добавить менеджера", callback_data="adm:addmanager")],
        [InlineKeyboardButton(text="🔄 Перезагрузить бота", callback_data="adm:restart")],
        [InlineKeyboardButton(text="✖️ Закрыть", callback_data="adm:exit")]]))


def load_prices():
    try:
        config = json.loads(PRICES_FILE.read_text("utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError("Не удалось прочитать prices.json.") from exc
    if not isinstance(config, dict):
        raise ValueError("Корень prices.json должен быть JSON-объектом.")
    if not isinstance(config.get("base"), dict):
        raise ValueError("В prices.json отсутствует объект base.")
    if not isinstance(config.get("range"), list) or len(config["range"]) != 2:
        raise ValueError("Поле range в prices.json должно содержать два коэффициента.")
    if not isinstance(config.get("rules"), list):
        raise ValueError("Поле rules в prices.json должно быть списком.")
    if not isinstance(config.get("currency"), str):
        raise ValueError("Поле currency в prices.json должно быть строкой.")

    for key, value in config["base"].items():
        if not isinstance(value, (int, float)) or isinstance(value, bool) or not math.isfinite(value) or value < 0:
            raise ValueError(f"Некорректная базовая цена «{key}» в prices.json.")
    for value in config["range"]:
        if not isinstance(value, (int, float)) or isinstance(value, bool) or not math.isfinite(value) or value <= 0:
            raise ValueError("Коэффициенты range должны быть положительными числами.")
    for index, rule in enumerate(config["rules"], 1):
        if not isinstance(rule, dict) or not isinstance(rule.get("step"), str):
            raise ValueError(f"Правило цен №{index} задано некорректно.")
        price_keys = [key for key in ("add", "per_item") if key in rule]
        if len(price_keys) != 1:
            raise ValueError(f"В правиле цен №{index} должно быть одно поле add или per_item.")
        if "add" in rule and not isinstance(rule.get("answer"), str):
            raise ValueError(f"В правиле цен №{index} отсутствует текст answer.")
        value = rule[price_keys[0]]
        if not isinstance(value, (int, float)) or isinstance(value, bool) or not math.isfinite(value) or value < 0:
            raise ValueError(f"Некорректная сумма в правиле цен №{index}.")
    return config


def price_fields(config):
    base_labels = {
        "p": "Товары",
        "s": "Услуги",
        "m": "Товары и услуги",
        "r": "Доработка сайта",
        "n": "Другое",
    }
    fields = [
        {"path": "currency", "label": "Валюта", "value": config["currency"], "kind": "text"},
        {"path": "range:0", "label": "Нижний коэффициент", "value": config["range"][0], "kind": "ratio"},
        {"path": "range:1", "label": "Верхний коэффициент", "value": config["range"][1], "kind": "ratio"},
    ]
    for code, value in config["base"].items():
        label = base_labels.get(code, code)
        fields.append({"path": f"base:{code}", "label": f"База · {label}",
                       "value": value, "kind": "amount"})
    for index, rule in enumerate(config["rules"]):
        key = "per_item" if "per_item" in rule else "add"
        detail = f"· {rule['answer']}" if "answer" in rule else "за пункт"
        label = f"{rule['step']} {detail} ({'за каждый выбор' if key == 'per_item' else 'доплата'})"
        fields.append({"path": f"rule:{index}:{key}", "label": label,
                       "value": rule[key], "kind": "amount"})
    return fields


def price_display(value):
    if not isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, float) and not value.is_integer():
        return f"{value:g}"
    return f"{int(value):,}".replace(",", " ")


def set_price_value(config, path, value):
    parts = path.split(":")
    if parts == ["currency"]:
        config["currency"] = value
    elif len(parts) == 2 and parts[0] == "range" and parts[1] in ("0", "1"):
        config["range"][int(parts[1])] = value
    elif len(parts) == 2 and parts[0] == "base" and parts[1] in config["base"]:
        config["base"][parts[1]] = value
    elif (len(parts) == 3 and parts[0] == "rule" and parts[1].isdigit()
          and 0 <= int(parts[1]) < len(config["rules"])
          and parts[2] in ("add", "per_item")
          and parts[2] in config["rules"][int(parts[1])]):
        config["rules"][int(parts[1])][parts[2]] = value
    else:
        raise ValueError("Выбранное поле цен больше не существует.")


async def show_prices(msg: Message, s):
    try:
        config = load_prices()
        fields = price_fields(config)
    except ValueError as exc:
        log.exception("Не удалось открыть редактор цен")
        await render(msg, s, f"⚠️ {exc}", InlineKeyboardMarkup(inline_keyboard=[
            btn("🔧 К панели управления", "adm:home")]))
        return

    text = "💰 Редактирование цен\n\nНажмите на значение, которое хотите изменить."
    buttons = [
        InlineKeyboardButton(
            text=f"{field['label'][:42]}: {price_display(field['value'])}"[:64],
            callback_data=f"adm:priceedit:{field['path']}",
        )
        for field in fields
    ]
    rows = [buttons[index:index + 2] for index in range(0, len(buttons), 2)]
    rows.append([InlineKeyboardButton(text="🔧 К панели управления", callback_data="adm:home")])
    await render(msg, s, text, InlineKeyboardMarkup(inline_keyboard=rows))


async def open_admin(m: Message, s):
    await remove_menu_editor_keyboard(m, s)
    s["ask"] = None
    s["ld_ask"] = False
    s["role_add"] = None
    s["price_edit"] = None
    s["menu_edit"] = None
    await clear_menu_hint(m, s)
    text, markup = admin_home()
    await render(m, s, text, markup)        # панель открывается в том же единственном меню


async def show_menu_editor(msg: Message, s, notice=""):
    editor = s.get("menu_edit")
    if not editor:
        editor = {"stack": ["main"], "action": "open"}
        s["menu_edit"] = editor
    sid = editor["stack"][-1]
    step = STEPS[sid]
    text = step_text(step).replace("{name}", client_name(s))
    if notice:
        await send_menu_hint(msg, s, notice)
    text = f"🛠 Редактор меню · {step['title']}\n\n{text}"

    rows = []
    for option in effective_options(step, include_hidden=True):
        label = f"🙈 {option['label']}" if option["_hidden"] else option["label"]
        rows.append(btn(label[:60], f"adm:menuitem:{option['_menu_key']}"))
    await render(msg, s, text, InlineKeyboardMarkup(inline_keyboard=rows))
    await show_menu_editor_keyboard(msg, s, step)


async def show_menu_editor_keyboard(msg: Message, s, step):
    rows = [
        [KeyboardButton(text="🧭 Открыть подменю"), KeyboardButton(text="💬 Текст вопроса")],
        [KeyboardButton(text="✏️ Изменить пункт"), KeyboardButton(text="👁 Скрыть / показать")],
    ]
    if allows_added_choices(step):
        rows.append([KeyboardButton(text="🗑 Удалить пункт"), KeyboardButton(text="➕ Добавить пункт")])
    else:
        rows.append([KeyboardButton(text="🗑 Удалить пункт")])
    next_step = step.get("next")
    if isinstance(next_step, str) and next_step in STEPS:
        rows.append([KeyboardButton(text="➡️ Следующий вопрос")])
    rows.extend([
        [KeyboardButton(text="↩️ Вернуть меню к стандартному виду")],
        [KeyboardButton(text="◀️ Предыдущий пункт"), KeyboardButton(text="🏠 В начало")],
        [KeyboardButton(text="🔧 К панели управления")],
    ])
    old_message_id = s.get("editor_keyboard_message")
    keyboard_message = await msg.bot.send_message(
        msg.chat.id,
        "\u2063",
        reply_markup=ReplyKeyboardMarkup(keyboard=rows, resize_keyboard=True),
    )
    s["editor_keyboard_message"] = keyboard_message.message_id
    asyncio.create_task(delete_keyboard_carrier(msg.bot, msg.chat.id, keyboard_message.message_id))
    if old_message_id:
        try:
            await msg.bot.delete_message(msg.chat.id, old_message_id)
        except TelegramBadRequest:
            log.debug("Не удалось удалить предыдущее сообщение управления редактором")


async def handle_menu_editor_keyboard(m: Message, s, action):
    editor = s.get("menu_edit")
    if not editor:
        return
    await try_delete(m)
    if action.startswith("menumode:"):
        mode = action.split(":", 1)[1]
        step = STEPS[editor["stack"][-1]]
        if mode == "add" and not allows_added_choices(step):
            await send_menu_hint(m, s, "Добавлять пункты можно только в выборах, которые записываются в заявку.")
            return
        if mode == "text":
            editor["pending"] = "step_text"
            await send_menu_hint(m, s, "Отправьте новый текст для этого пункта меню.")
            return
        if mode == "add":
            editor["pending"] = "add_label"
            await send_menu_hint(m, s, "Отправьте текст нового пункта меню.")
            return
        editor["action"] = mode
        labels = {
            "open": "Выберите пункт, чтобы открыть его следующее меню.",
            "edit": "Выберите пункт, текст которого нужно изменить.",
            "hide": "Выберите пункт, который нужно скрыть или показать.",
            "delete": "Выберите пункт, который нужно удалить.",
        }
        await show_menu_editor(m, s, labels.get(mode, "Выберите пункт меню."))
        return

    await clear_menu_hint(m, s)
    if action == "menunext":
        step = STEPS[editor["stack"][-1]]
        next_step = step.get("next")
        if isinstance(next_step, str) and next_step in STEPS:
            push(editor, next_step)
            editor["action"] = "open"
            editor.pop("pending", None)
            await show_menu_editor(m, s)
        return
    if action == "menuprev":
        editor.pop("pending", None)
        editor.pop("choice_key", None)
        editor.pop("pending_label", None)
        editor.pop("routes", None)
        if len(editor["stack"]) > 1:
            editor["stack"].pop()
            editor["action"] = "open"
            await show_menu_editor(m, s)
        else:
            await send_menu_hint(m, s, "Вы уже в первом меню.")
        return
    if action == "menustart":
        editor["stack"] = ["main"]
        editor["action"] = "open"
        editor.pop("pending", None)
        editor.pop("choice_key", None)
        editor.pop("pending_label", None)
        editor.pop("routes", None)
        await show_menu_editor(m, s)
        return
    if action == "menureset":
        sid = editor["stack"][-1]
        overrides.pop(sid, None)
        settings.get("menu_choices", {}).pop(sid, None)
        if not settings.get("menu_choices"):
            settings.pop("menu_choices", None)
        save_json(OVERRIDES_FILE, overrides)
        save_json(SETTINGS_FILE, settings)
        editor["action"] = "open"
        editor.pop("pending", None)
        editor.pop("choice_key", None)
        editor.pop("pending_label", None)
        editor.pop("routes", None)
        await show_menu_editor(m, s)
        await send_menu_hint(m, s, "✅ Этот экран и его пункты возвращены к стандартному виду.")
        return
    if action == "home":
        s["menu_edit"] = None
        await remove_menu_editor_keyboard(m, s)
        text, markup = admin_home()
        await render(m, s, text, markup)


def save_menu_choices():
    save_json(SETTINGS_FILE, settings)


async def request_restart(user_id):
    global restart_requested
    settings["restart_menu_user"] = user_id
    save_json(SETTINGS_FILE, settings)
    restart_requested = True
    await runtime_dispatcher.stop_polling()


@router.message(Command("admin"))
async def cmd_admin(m: Message):
    s = sess(m.from_user)
    if not await require_privacy(m, s):
        return
    await clear_chat(m, s)
    if not is_admin(m.from_user.id):
        await answer_temp(m, f"Нет доступа. Ваш ID: {m.from_user.id}.\nЧтобы стать администратором, впишите "
                             "этот ID в ADMIN_IDS в файле .env или попросите текущего администратора добавить "
                             "вас через админ-панель.", delay=60)
        return
    await open_admin(m, s)


@router.message(Command("setmanager"))
async def cmd_setmanager(m: Message):
    s = sess(m.from_user)
    if not await require_privacy(m, s):
        return
    await clear_chat(m, s)
    if not is_admin(m.from_user.id):
        return
    settings["manager_chat"] = m.chat.id
    save_json(SETTINGS_FILE, settings)
    await answer_temp(m, "✅ Готово. Новые заявки будут приходить в этот чат.")
    await update_menu(m.bot, m.chat.id, sess(m.from_user), force=True)   # в меню появится «Последние заявки»


@router.message(Command("contacts"))
async def cmd_contacts(m: Message):
    s = sess(m.from_user)
    await clear_chat(m, s, preserve_current=True)
    if not privacy_confirmed(m.from_user.id):
        s.update(stack=[], ans={}, files=[], ask=None, submitted=False)
        await show_privacy(m, s)
        await try_delete(m)
        return
    s["ask"] = None
    s["role_add"] = None
    s["menu_edit"] = None
    if not s["stack"]:
        s["stack"] = ["main"]
    await go(m, s, "contacts")
    await try_delete(m)


@router.message(Command("leads"))
async def cmd_leads(m: Message):
    s = sess(m.from_user)
    if not await require_privacy(m, s):
        return
    await clear_chat(m, s)
    if not (is_admin(m.from_user.id) or is_manager(m.from_user.id)):
        return
    s["ask"] = None
    s["role_add"] = None
    s["menu_edit"] = None
    await show_leads(m, s, 0)


def btn(text, data):
    return [InlineKeyboardButton(text=text, callback_data=data)]


@router.callback_query(F.data.startswith("adm:"))
async def cb_admin(cb: CallbackQuery):
    s = sess(cb.from_user)
    if not await require_privacy_callback(cb, s):
        return
    if not is_admin(cb.from_user.id):
        await cb.answer("Нет доступа", show_alert=True)
        return
    s["menu"] = cb.message.message_id
    _, act, *rest = cb.data.split(":", 2)
    arg = rest[0] if rest else ""
    groups = list(GROUPS)
    s["adm_edit"] = None
    s["role_add"] = None
    s["price_edit"] = None
    if act == "section":
        await cb.answer()
        return
    elif act == "menu":
        s["ask"] = None
        s["ld_ask"] = False
        s["menu_edit"] = {"stack": ["main"], "action": "open"}
        await clear_menu_hint(cb.message, s)
        await show_menu_editor(cb.message, s)
        await cb.answer()
        return
    elif act == "menunext":
        editor = s.get("menu_edit")
        if not editor:
            await cb.answer("Редактор завершён. Откройте его заново.", show_alert=True)
            return
        step = STEPS[editor["stack"][-1]]
        next_step = step.get("next")
        if not isinstance(next_step, str) or next_step not in STEPS:
            await cb.answer("У этого вопроса нет следующего шага.", show_alert=True)
            return
        push(editor, next_step)
        editor["action"] = "open"
        editor.pop("pending", None)
        await clear_menu_hint(cb.message, s)
        await show_menu_editor(cb.message, s)
        await cb.answer()
        return
    elif act == "menuprev":
        editor = s.get("menu_edit")
        if not editor:
            await cb.answer("Редактор завершён. Откройте его заново.", show_alert=True)
            return
        await clear_menu_hint(cb.message, s)
        editor.pop("pending", None)
        editor.pop("choice_key", None)
        editor.pop("pending_label", None)
        editor.pop("routes", None)
        if len(editor["stack"]) > 1:
            editor["stack"].pop()
            editor["action"] = "open"
            await show_menu_editor(cb.message, s)
            await cb.answer()
        else:
            await cb.answer("Вы уже в первом меню.", show_alert=True)
        return
    elif act == "menustart":
        editor = s.get("menu_edit")
        if not editor:
            await cb.answer("Редактор завершён. Откройте его заново.", show_alert=True)
            return
        await clear_menu_hint(cb.message, s)
        editor["stack"] = ["main"]
        editor["action"] = "open"
        editor.pop("pending", None)
        editor.pop("choice_key", None)
        editor.pop("pending_label", None)
        editor.pop("routes", None)
        await show_menu_editor(cb.message, s)
        await cb.answer()
        return
    elif act == "prices":
        await clear_menu_hint(cb.message, s)
        await show_prices(cb.message, s)
        await cb.answer()
        return
    elif act == "priceedit":
        try:
            config = load_prices()
            field = next((item for item in price_fields(config) if item["path"] == arg), None)
            if field is None:
                raise ValueError("Это поле цен больше не существует.")
        except ValueError as exc:
            await cb.answer(str(exc), show_alert=True)
            await show_prices(cb.message, s)
            return
        s["price_edit"] = {"path": arg}
        prompt_text = ("Отправьте новое значение валюты." if field["kind"] == "text"
                       else f"Отправьте новое значение для «{field['label']}».\n"
                            + ("Допустима десятичная дробь, например 0.85." if field["kind"] == "ratio"
                               else "Введите целое число без разделителей."))
        await send_menu_hint(cb.message, s, prompt_text, InlineKeyboardMarkup(inline_keyboard=[
            btn("Отмена", "adm:pricecancel")]))
        await cb.answer()
        return
    elif act == "pricecancel":
        s["price_edit"] = None
        await clear_menu_hint(cb.message, s)
        await show_prices(cb.message, s)
        await cb.answer()
        return
    elif act == "resetall":
        text = "Сбросить все тексты и настройки пунктов меню к стандартному виду? Это действие нельзя отменить."
        markup = InlineKeyboardMarkup(inline_keyboard=[
            btn("✅ Да, сбросить всё", "adm:resetallconfirm"),
            btn("Отмена", "adm:home"),
        ])
        await render(cb.message, s, text, markup)
        await cb.answer()
        return
    elif act == "resetallconfirm":
        overrides.clear()
        settings.pop("menu_choices", None)
        save_json(OVERRIDES_FILE, overrides)
        save_json(SETTINGS_FILE, settings)
        s["menu_edit"] = None
        await clear_menu_hint(cb.message, s)
        text, markup = admin_home()
        await render(cb.message, s, text, markup)
        await send_menu_hint(cb.message, s, "✅ Все меню и выборы возвращены к стандартному виду.")
        await cb.answer("Все меню сброшены.")
        return
    elif act == "menureset":
        editor = s.get("menu_edit")
        if not editor:
            await cb.answer("Редактор завершён. Откройте его заново.", show_alert=True)
            return
        sid = editor["stack"][-1]
        overrides.pop(sid, None)
        settings.get("menu_choices", {}).pop(sid, None)
        if not settings.get("menu_choices"):
            settings.pop("menu_choices", None)
        save_json(OVERRIDES_FILE, overrides)
        save_json(SETTINGS_FILE, settings)
        editor["action"] = "open"
        editor.pop("pending", None)
        editor.pop("choice_key", None)
        editor.pop("pending_label", None)
        editor.pop("routes", None)
        await clear_menu_hint(cb.message, s)
        await show_menu_editor(cb.message, s)
        await send_menu_hint(cb.message, s, "✅ Этот экран и его пункты возвращены к стандартному виду.")
        await cb.answer("Текущий экран сброшен.")
        return
    elif act == "menuback":
        editor = s.get("menu_edit")
        if editor and editor.get("pending"):
            await clear_menu_hint(cb.message, s)
            editor.pop("pending", None)
            editor.pop("choice_key", None)
            editor.pop("pending_label", None)
            editor.pop("routes", None)
            editor["action"] = "open"
            await show_menu_editor(cb.message, s)
        elif editor and len(editor["stack"]) > 1:
            await clear_menu_hint(cb.message, s)
            editor["stack"].pop()
            editor["action"] = "open"
            await show_menu_editor(cb.message, s)
        else:
            await clear_menu_hint(cb.message, s)
            s["menu_edit"] = None
            await remove_menu_editor_keyboard(cb.message, s)
            text, markup = admin_home()
            await render(cb.message, s, text, markup)
        await cb.answer()
        return
    elif act == "menumode":
        editor = s.get("menu_edit")
        if not editor:
            await cb.answer("Редактор завершён. Откройте его заново.", show_alert=True)
            return
        mode = arg
        if mode not in ("open", "text", "edit", "hide", "delete", "add"):
            await cb.answer("Неизвестное действие.", show_alert=True)
            return
        if mode == "add" and not allows_added_choices(STEPS[editor["stack"][-1]]):
            editor["action"] = "open"
            await show_menu_editor(cb.message, s)
            await cb.answer("Добавлять пункты можно только в выборах, которые записываются в заявку.",
                            show_alert=True)
            return
        if mode == "text":
            editor["pending"] = "step_text"
            await send_menu_hint(cb.message, s, "Отправьте новый текст для этого пункта меню.")
        elif mode == "add":
            editor["pending"] = "add_label"
            await send_menu_hint(cb.message, s, "Отправьте текст нового пункта меню.")
        else:
            editor["action"] = mode
            labels = {
                "open": "Выберите пункт, чтобы открыть его следующее меню.",
                "edit": "Выберите пункт, текст которого нужно изменить.",
                "hide": "Выберите пункт, который нужно скрыть или показать.",
                "delete": "Выберите пункт, который нужно удалить.",
            }
            await show_menu_editor(cb.message, s, labels.get(mode, "Выберите пункт меню."))
        await cb.answer()
        return
    elif act == "menuitem":
        editor = s.get("menu_edit")
        if not editor:
            await cb.answer("Редактор завершён. Откройте его заново.", show_alert=True)
            return
        sid = editor["stack"][-1]
        step = STEPS[sid]
        options = effective_options(step, include_hidden=True)
        if arg.isdigit():
            index = int(arg)
            option = options[index] if 0 <= index < len(options) else None
        else:
            option = next((item for item in options if item["_menu_key"] == arg), None)
        if option is None:
            await cb.answer("Пункт меню уже изменился. Обновите редактор.", show_alert=True)
            return
        key = option["_menu_key"]
        action = editor.get("action", "open")
        if action == "open":
            if option["_hidden"]:
                await cb.answer("Сначала покажите скрытый пункт.", show_alert=True)
                return
            target = next_of(step, option, s)
            if target not in STEPS:
                await cb.answer("У этого пункта нет подменю.", show_alert=True)
                return
            editor["stack"].append(target)
            editor["action"] = "open"
            await clear_menu_hint(cb.message, s)
            await show_menu_editor(cb.message, s)
        elif action == "edit":
            editor["pending"] = "choice_label"
            editor["choice_key"] = key
            await send_menu_hint(cb.message, s, f"Отправьте новый текст пункта:\n\n{option['label']}")
        elif action == "hide":
            config = menu_choice_config(step, create=True)
            hidden = config["hidden"]
            if key in hidden:
                hidden.remove(key)
                notice = f"Пункт «{option['label']}» снова показан."
            else:
                hidden.append(key)
                notice = f"Пункт «{option['label']}» скрыт от пользователей."
            save_menu_choices()
            editor["action"] = "open"
            await clear_menu_hint(cb.message, s)
            await show_menu_editor(cb.message, s)
            await cb.answer(notice)
            return
        elif action == "delete":
            target = option.get("next") or step.get("next")
            has_submenu = (
                option.get("next") in STEPS
                or target in STEPS and (
                    not option.get("rec", True)
                    or not step.get("record", True)
                    or not STEPS[target].get("record", True)
                )
            )
            if has_submenu:
                editor["action"] = "open"
                await show_menu_editor(cb.message, s,
                                       "Этот пункт ведёт в подменю, поэтому его можно только скрыть.")
                await cb.answer()
                return
            deleted = menu_choice_config(step, create=True)["deleted"]
            if key not in deleted:
                deleted.append(key)
            save_menu_choices()
            editor["action"] = "open"
            await clear_menu_hint(cb.message, s)
            await show_menu_editor(cb.message, s)
            await cb.answer("Пункт удалён.")
            return
        await cb.answer()
        return
    elif act == "menuroute":
        editor = s.get("menu_edit")
        routes = editor.get("routes", []) if editor else []
        index = int(arg)
        if not 0 <= index < len(routes):
            await cb.answer("Маршрут устарел. Попробуйте добавить пункт снова.", show_alert=True)
            return
        route = routes[index]
        sid = editor["stack"][-1]
        config = menu_choice_config(STEPS[sid], create=True)
        added = config["added"]
        new_index = max([int(item["id"].split("_", 1)[1]) for item in added
                         if item["id"].startswith("added_") and item["id"].split("_", 1)[1].isdigit()]
                        + [-1]) + 1
        added.append({"id": f"added_{new_index}", "label": editor.pop("pending_label"),
                      "next": route["next"], "rec": route["rec"], "exclusive": False,
                      "ask": None, "say": None})
        editor.pop("routes", None)
        editor.pop("pending", None)
        save_menu_choices()
        editor["action"] = "open"
        await clear_menu_hint(cb.message, s)
        await show_menu_editor(cb.message, s)
        await cb.answer("Новый пункт добавлен.")
        return
    elif act == "exit":                       # закрыть панель: вернуться на прежний экран
        s["menu_edit"] = None
        await clear_menu_hint(cb.message, s)
        if s["stack"]:
            await show(cb.message, s, s["stack"][-1])
        else:
            await go(cb.message, s, "main")
        await cb.answer()
        return
    if act == "home":
        await clear_menu_hint(cb.message, s)
        s["menu_edit"] = None
        await remove_menu_editor_keyboard(cb.message, s)
        text, markup = admin_home()
    elif act == "groups":
        text = "Выберите раздел сценария:"
        markup = InlineKeyboardMarkup(inline_keyboard=[btn(g, f"adm:g:{i}") for i, g in enumerate(groups)]
                                      + [btn(BACK, "adm:home")])
    elif act == "g":
        g = groups[int(arg)]
        text = f"{g}\nВыберите сообщение для редактирования (✏️ — уже изменено вами):"
        rows = [btn(("✏️ " if sid in overrides else "") + STEPS[sid]["title"][:55], f"adm:s:{sid}")
                for sid in GROUPS[g]]
        markup = InlineKeyboardMarkup(inline_keyboard=rows + [btn(BACK, "adm:groups")])
    elif act in ("s", "r"):
        if act == "r":
            overrides.pop(arg, None)
            save_json(OVERRIDES_FILE, overrides)
        st = STEPS[arg]
        text = f"«{st['title']}»\n\nТекущий текст:\n\n{step_text(st)[:3500]}"
        rows = [btn("✏️ Изменить текст", f"adm:e:{arg}")]
        if arg in overrides:
            rows.append(btn("↩️ Вернуть исходный текст", f"adm:r:{arg}"))
        rows.append(btn(BACK, f"adm:g:{list(GROUPS).index(st['group'])}"))
        markup = InlineKeyboardMarkup(inline_keyboard=rows)
    elif act == "e":
        s["adm_edit"] = arg
        text = ("Отправьте новый текст одним сообщением.\n"
                + ("Можно использовать {name} — туда подставится имя клиента.\n" if arg == "f_done" else "")
                + "Эмодзи и переносы строк сохранятся.")
        markup = InlineKeyboardMarkup(inline_keyboard=[btn("Отмена", f"adm:s:{arg}")])
    elif act == "mgr":
        cur = manager_chat()
        text = (f"Заявки приходят в чат: {cur or 'не задан'}.\n\nЧтобы заявки приходили в другой чат "
                "(личный, группа), напишите команду /setmanager в нужном чате.\n"
                "Все заявки также сохраняются в файл data/leads.csv (открывается в Excel).")
        markup = InlineKeyboardMarkup(inline_keyboard=[btn(BACK, "adm:home")])
    elif act in ("addadmin", "addmanager"):
        role = "admin" if act == "addadmin" else "manager"
        s["role_add"] = role
        role_name = "администраторов" if role == "admin" else "менеджеров"
        text = (f"Отправьте Telegram ID {role_name} одним сообщением.\n"
                "Можно указать несколько ID через запятую или пробел. Узнать ID можно командой /myid.")
        markup = InlineKeyboardMarkup(inline_keyboard=[btn("Отмена", "adm:home")])
    elif act == "restart":
        await cb.message.edit_text("🔄 Перезагружаю бота…")
        await cb.answer()
        await request_restart(cb.from_user.id)
        return
    else:
        await cb.answer()
        return
    try:
        await cb.message.edit_text(text, reply_markup=markup)
    except TelegramBadRequest:
        pass
    await cb.answer()


# ───────────────────────── Меню заявок ─────────────────────────
def leads_screen(page, notice=""):
    leads = read_leads()[::-1]                      # свежие сверху
    pages = max(1, -(-len(leads) // LEADS_PAGE))
    page = min(max(page, 0), pages - 1)
    if leads:
        text = (f"📋 Заявки: всего {len(leads)}, стр. {page + 1}/{pages}\n"
                "Нажмите на заявку, чтобы открыть подробности.")
    else:
        text = "📋 Заявок пока нет."
    if notice:
        text = f"{notice}\n\n{text}"
    rows = [btn(f"№{l['ID']} · {l['Дата'][8:10]}.{l['Дата'][5:7]} · {l['Имя'][:16]} · {l['Ветка'][:22]}",
                f"ld:v:{l['ID']}") for l in leads[page * LEADS_PAGE:(page + 1) * LEADS_PAGE]]
    nav = []
    if page > 0:
        nav.append(InlineKeyboardButton(text="◀️", callback_data=f"ld:l:{page - 1}"))
    if page < pages - 1:
        nav.append(InlineKeyboardButton(text="▶️", callback_data=f"ld:l:{page + 1}"))
    if nav:
        rows.append(nav)
    if leads:
        rows += [btn("📥 Скачать все в Excel", "ld:xlsx"), btn("🗑 Удалить по ID", "ld:ask")]
    rows.append(btn("✖️ Закрыть", "ld:close"))
    return text, InlineKeyboardMarkup(inline_keyboard=rows), page


async def show_leads(msg: Message, s, page=0, notice=""):
    s.update(ld_ask=False, ask=None)
    text, markup, s["ld_page"] = leads_screen(page, notice)
    await render(msg, s, text, markup)


async def confirm_delete(msg: Message, s, ids, notice=""):
    s["ld_pending"] = ids
    text = f"Удалить заявки {', '.join('№' + i for i in ids)}? Это нельзя отменить."
    if notice:
        text = f"{notice}\n\n{text}"
    await render(msg, s, text, InlineKeyboardMarkup(inline_keyboard=[
        btn("✅ Да, удалить", "ld:dok"), btn("Отмена", f"ld:l:{s.get('ld_page', 0)}")]))


async def ask_lead_ids(msg: Message, s, notice=""):
    s["ld_ask"] = True
    text = ("Отправьте ID заявки одним сообщением (он в первой колонке Excel-файла и в карточке заявки).\n"
            "Можно несколько: 3, 5, 8.")
    if notice:
        text = f"{notice}\n\n{text}"
    await render(msg, s, text, InlineKeyboardMarkup(inline_keyboard=[btn("Отмена", f"ld:l:{s.get('ld_page', 0)}")]))


@router.callback_query(F.data.startswith("ld:"))
async def cb_leads(cb: CallbackQuery):
    s = sess(cb.from_user)
    if not await require_privacy_callback(cb, s):
        return
    uid = cb.from_user.id
    if not (is_admin(uid) or is_manager(uid)):
        await cb.answer("Нет доступа", show_alert=True)
        return
    s["menu"] = cb.message.message_id
    s["menu_edit"] = None
    await clear_menu_hint(cb.message, s)
    s["ld_ask"] = False
    _, act, *rest = cb.data.split(":")
    arg = rest[0] if rest else ""

    if act == "close":
        if s["stack"]:
            await show(cb.message, s, s["stack"][-1])
        else:
            await go(cb.message, s, "main")
    elif act == "l":
        await show_leads(cb.message, s, int(arg))
    elif act == "v":
        lead = next((l for l in read_leads() if l["ID"] == arg), None)
        if not lead:
            await show_leads(cb.message, s, s.get("ld_page", 0), "⚠️ Заявка не найдена (возможно, уже удалена).")
        else:
            await render(cb.message, s, lead_card(lead), InlineKeyboardMarkup(inline_keyboard=[
                btn("🗑 Удалить эту заявку", f"ld:del:{arg}"), btn("◀️ К списку", f"ld:l:{s.get('ld_page', 0)}")]))
    elif act == "del":
        await confirm_delete(cb.message, s, [arg])
    elif act == "ask":
        await ask_lead_ids(cb.message, s)
    elif act == "dok":
        ids = set(s.pop("ld_pending", []))
        leads = read_leads()
        keep = [l for l in leads if l["ID"] not in ids]
        write_leads(keep)
        await show_leads(cb.message, s, s.get("ld_page", 0), f"🗑 Удалено заявок: {len(leads) - len(keep)}.")
    elif act == "xlsx":
        leads = read_leads()
        if not leads:
            await cb.answer("Заявок пока нет", show_alert=True)
            return
        stamp = datetime.now().strftime("%Y%m%d_%H%M")
        try:
            doc = BufferedInputFile(build_xlsx(leads), filename=f"leads_{stamp}.xlsx")
            await cb.message.answer_document(doc, caption=f"📥 Все заявки ({len(leads)} шт.)")
        except ImportError:
            await cb.message.answer_document(FSInputFile(LEADS_FILE, filename=f"leads_{stamp}.csv"),
                                             caption="Библиотека openpyxl не установлена (pip install openpyxl), "
                                                     "поэтому отправлен CSV — он тоже открывается в Excel.")
        await cb.answer("Файл отправлен")
        return
    await cb.answer()


# ───────────────────────── Кнопки сценария ─────────────────────────
async def guard(cb: CallbackQuery):
    s = sess(cb.from_user)
    if not await require_privacy_callback(cb, s):
        return None
    if not s["stack"]:
        await cb.answer("Сессия устарела, отправьте /start", show_alert=True)
        return None
    s["menu"] = cb.message.message_id
    return s


@router.callback_query(F.data == "back")
async def cb_back(cb: CallbackQuery):
    if not (s := await guard(cb)):
        return
    s["ask"] = None
    if len(s["stack"]) > 1:
        s["stack"].pop()
    await show(cb.message, s, s["stack"][-1])
    await cb.answer()


@router.callback_query(F.data == "cancelask")
async def cb_cancelask(cb: CallbackQuery):
    if not (s := await guard(cb)):
        return
    ask, s["ask"] = s["ask"], None
    cur = STEPS[s["stack"][-1]]
    if ask and cur["type"] == "multi":      # отмена ввода в мульти-выборе снимает отметку
        a = s["ans"].get(ask["step"])
        if a:
            label = ask["option"]["label"]
            if label in a["sel"]:
                a["sel"].remove(label)
            a["extra"].pop(label, None)
    await show(cb.message, s, cur["id"])
    await cb.answer()


@router.callback_query(F.data == "done")
async def cb_done(cb: CallbackQuery):
    if not (s := await guard(cb)):
        return
    s["ask"] = None
    step = STEPS[s["stack"][-1]]
    await go(cb.message, s, next_of(step, None, s))
    await cb.answer()


@router.callback_query(F.data.startswith("o:"))
async def cb_option(cb: CallbackQuery):
    if not (s := await guard(cb)):
        return
    step = STEPS[s["stack"][-1]]
    options = effective_options(step)
    option_id = cb.data[2:]
    if option_id.isdigit():
        legacy_key = f"base_{int(option_id)}"
        opt = next((option for option in effective_options(step, include_hidden=True)
                    if option["_menu_key"] == legacy_key), None)
    else:
        opt = next((option for option in options if option["_menu_key"] == option_id), None)
    if opt is None:
        await cb.answer("Этот пункт больше недоступен.")
        return
    if opt["_hidden"]:
        await cb.answer("Этот пункт скрыт.")
        return
    sid = step["id"]
    label = opt["label"]

    if step["type"] == "multi":
        a = s["ans"].setdefault(sid, {"title": step["title"], "sel": [], "extra": {}})
        excl = {x["label"] for x in options if x["exclusive"]}
        if label in a["sel"]:
            a["sel"].remove(label)
            a["extra"].pop(label, None)
        else:
            if opt["exclusive"]:
                a["sel"].clear()
                a["extra"].clear()
            else:
                a["sel"] = [x for x in a["sel"] if x not in excl]
                for x in excl:
                    a["extra"].pop(x, None)
            a["sel"].append(label)
        s["ask"] = {"step": sid, "option": opt} if opt["ask"] and label in a["sel"] else None
        if s["ask"]:
            await prompt(cb.message, s, opt["ask"])
        else:
            await show(cb.message, s, sid)
        await cb.answer()
        return

    s["ask"] = None
    if step["record"] and opt["rec"]:
        for pref in step.get("purge", ()):      
            for k in [k for k in s["ans"] if k.startswith(pref)]:
                del s["ans"][k]
        s["ans"][sid] = {"title": step["title"], "sel": [label], "extra": {}}
    if opt["ask"]:
        s["ask"] = {"step": sid, "option": opt}
        await prompt(cb.message, s, opt["ask"])
        await cb.answer()
        return
    await go(cb.message, s, next_of(step, opt, s), notice=opt["say"] or "")
    await cb.answer()


# ───────────────────────── Текстовые сообщения ─────────────────────────
def extract(m: Message, s):
    """Достаёт текст/файл из сообщения клиента."""
    note = m.caption or ""
    if m.text:
        return m.text.strip()
    if m.document:
        s["files"].append(("document", m.document.file_id))
        return f"[файл: {m.document.file_name}] {note}".strip()
    if m.photo:
        s["files"].append(("photo", m.photo[-1].file_id))
        return f"[фото] {note}".strip()
    return None


@router.message()
async def on_message(m: Message):
    s = sess(m.from_user)
    if not privacy_confirmed(m.from_user.id):
        await clear_chat(m, s)
        await show_privacy(m, s)
        return

    editor = s.get("menu_edit")
    if editor and is_admin(m.from_user.id) and m.text:
        editor_actions = {
            "🧭 Открыть подменю": "menumode:open",
            "💬 Текст вопроса": "menumode:text",
            "✏️ Изменить пункт": "menumode:edit",
            "👁 Скрыть / показать": "menumode:hide",
            "🗑 Удалить пункт": "menumode:delete",
            "➕ Добавить пункт": "menumode:add",
            "➡️ Следующий вопрос": "menunext",
            "↩️ Вернуть меню к стандартному виду": "menureset",
            "◀️ Предыдущий пункт": "menuprev",
            "🏠 В начало": "menustart",
            "🔧 К панели управления": "home",
        }
        action = editor_actions.get(m.text.strip())
        if action:
            await handle_menu_editor_keyboard(m, s, action)
            return

    price_edit = s.get("price_edit")
    if price_edit and is_admin(m.from_user.id):
        if not m.text or not m.text.strip():
            await try_delete(m)
            await send_menu_hint(m, s, "Отправьте новое значение текстом.", InlineKeyboardMarkup(
                inline_keyboard=[btn("Отмена", "adm:pricecancel")]))
            return
        try:
            config = load_prices()
            field = next((item for item in price_fields(config)
                          if item["path"] == price_edit["path"]), None)
            if field is None:
                raise ValueError("Это поле цен больше не существует.")
            value_text = m.text.strip()
            if field["kind"] == "text":
                if len(value_text) > 12 or "\n" in value_text:
                    raise ValueError("Название валюты должно содержать от 1 до 12 символов.")
                value = value_text
            else:
                value = float(value_text.replace(",", "."))
                if not math.isfinite(value):
                    raise ValueError("Введите конечное число.")
                if field["kind"] == "ratio":
                    if value <= 0:
                        raise ValueError("Коэффициент должен быть больше нуля.")
                    range_index = int(field["path"].split(":")[1])
                    other_value = config["range"][1 - range_index]
                    if (range_index == 0 and value > other_value) or (range_index == 1 and value < other_value):
                        raise ValueError("Нижний коэффициент не может быть больше верхнего.")
                else:
                    if value < 0 or not value.is_integer():
                        raise ValueError("Цена должна быть целым неотрицательным числом.")
                    value = int(value)
            set_price_value(config, price_edit["path"], value)
            save_json(PRICES_FILE, config)
        except ValueError as exc:
            await try_delete(m)
            await send_menu_hint(m, s, f"⚠️ {exc}\nПопробуйте ещё раз.", InlineKeyboardMarkup(
                inline_keyboard=[btn("Отмена", "adm:pricecancel")]))
            return
        except OSError:
            log.exception("Не удалось сохранить prices.json")
            await try_delete(m)
            await send_menu_hint(m, s, "⚠️ Не удалось сохранить prices.json. Проверьте доступ к файлу.",
                                 InlineKeyboardMarkup(inline_keyboard=[btn("Отмена", "adm:pricecancel")]))
            return
        s["price_edit"] = None
        await try_delete(m)
        await clear_menu_hint(m, s)
        await show_prices(m, s)
        return

    if not is_admin(m.from_user.id):
        current_step = STEPS.get(s["stack"][-1]) if s["stack"] else None
        if not s.get("ask") and not (current_step and current_step.get("accept_text")):
            await try_delete(m)
            return

    if editor and is_admin(m.from_user.id) and editor.get("pending") == "route":
        await try_delete(m)
        return
    if editor and editor.get("pending") and is_admin(m.from_user.id):
        if not m.text or not m.text.strip():
            await try_delete(m)
            await send_menu_hint(m, s, "Отправьте непустой текст.")
            return
        value = m.text.strip()
        if len(value) > 64 and editor["pending"] in ("choice_label", "add_label"):
            await try_delete(m)
            await send_menu_hint(m, s, "Текст пункта должен быть не длиннее 64 символов.")
            return
        pending = editor.pop("pending")
        await try_delete(m)
        if pending == "step_text":
            overrides[editor["stack"][-1]] = m.text
            save_json(OVERRIDES_FILE, overrides)
            await clear_menu_hint(m, s)
            await show_menu_editor(m, s)
            return
        if pending == "choice_label":
            sid = editor["stack"][-1]
            config = menu_choice_config(STEPS[sid], create=True)
            config["labels"][editor.pop("choice_key")] = value
            save_menu_choices()
            await clear_menu_hint(m, s)
            await show_menu_editor(m, s)
            return
        if pending == "add_label":
            sid = editor["stack"][-1]
            step = STEPS[sid]
            editor["pending_label"] = value
            editor["pending"] = "route"
            default_next = step.get("next")
            if isinstance(default_next, str) and default_next in STEPS:
                editor["routes"] = [{"label": "Следующий шаг вопроса", "next": default_next,
                                     "rec": step.get("record", True)}]
            else:
                routes = []
                for option in effective_options(step, include_hidden=True):
                    target = next_of(step, option, s)
                    if isinstance(target, str) and target in STEPS:
                        routes.append({"label": option["label"], "next": target, "rec": option["rec"]})
                editor["routes"] = routes
            if not editor["routes"]:
                editor.pop("pending_label", None)
                await clear_menu_hint(m, s)
                await show_menu_editor(m, s, "Нельзя добавить пункт: для этого меню не настроен следующий шаг.")
                return
            route_buttons = [btn(route["label"][:55], f"adm:menuroute:{i}")
                             for i, route in enumerate(editor["routes"])]
            await send_menu_hint(m, s, "Выберите, куда будет вести новый пункт:",
                                 InlineKeyboardMarkup(inline_keyboard=route_buttons))
            return

    if editor and is_admin(m.from_user.id):
        await try_delete(m)
        return

    if s.get("role_add") and is_admin(m.from_user.id):
        role = s["role_add"]
        if not m.text:
            await render(m, s, "Отправьте Telegram ID текстом.",
                         InlineKeyboardMarkup(inline_keyboard=[btn("Отмена", "adm:home")]))
            return
        if not re.fullmatch(r"\d+(?:[\s,]+\d+)*", m.text.strip()):
            await try_delete(m)
            await render(m, s, "⚠️ Укажите один или несколько ID цифрами через запятую или пробел.",
                         InlineKeyboardMarkup(inline_keyboard=[btn("Отмена", "adm:home")]))
            return
        ids = list(dict.fromkeys(str(int(value)) for value in re.findall(r"\d+", m.text)))
        if any(value == "0" for value in ids):
            await try_delete(m)
            await render(m, s, "⚠️ Telegram ID должен быть положительным числом.",
                         InlineKeyboardMarkup(inline_keyboard=[btn("Отмена", "adm:home")]))
            return
        setting_key = "admin_ids" if role == "admin" else "manager_ids"
        current = settings.get(setting_key, [])
        if not isinstance(current, list):
            current = []
        current = list(dict.fromkeys(str(value) for value in current))
        new_ids = [value for value in ids if value not in current]
        s["role_add"] = None
        await try_delete(m)
        if not new_ids:
            await render(m, s, "ℹ️ Все указанные ID уже добавлены.",
                         InlineKeyboardMarkup(inline_keyboard=[btn("🔧 К админ-панели", "adm:home")]))
            return
        settings[setting_key] = current + new_ids
        save_json(SETTINGS_FILE, settings)
        role_name = "администраторы" if role == "admin" else "менеджеры"
        await render(m, s, f"✅ Добавлено в список ({role_name}): {', '.join(new_ids)}.\n"
                           "🔄 Перезагружаю бота…",
                     InlineKeyboardMarkup(inline_keyboard=[btn("🔧 К админ-панели", "adm:home")]))
        await request_restart(m.from_user.id)
        return

    if s.get("adm_edit") and is_admin(m.from_user.id) and m.text:       # администратор правит текст
        sid, s["adm_edit"] = s["adm_edit"], None
        overrides[sid] = m.text
        save_json(OVERRIDES_FILE, overrides)
        await try_delete(m)
        await render(m, s, "✅ Текст сохранён. Бот уже использует новый вариант.",
                     InlineKeyboardMarkup(inline_keyboard=[btn("🔧 К админ-панели", "adm:home")]))
        return

    if s.get("ld_ask") and m.text and (is_admin(m.from_user.id) or is_manager(m.from_user.id)):
        await try_delete(m)                         # администратор вводит ID заявок для удаления
        ids = list(dict.fromkeys(str(int(x)) for x in re.findall(r"\d+", m.text)))
        have = {l["ID"] for l in read_leads()}
        found, missing = [i for i in ids if i in have], [i for i in ids if i not in have]
        note = f"⚠️ Не найдены ID: {', '.join(missing)}." if missing else ""
        if found:
            s["ld_ask"] = False
            await confirm_delete(m, s, found, note)
        else:
            await ask_lead_ids(m, s, note or "⚠️ Не нашёл в сообщении ни одного ID.")
        return

    if not s["stack"]:
        await cmd_start(m)
        return
    step = STEPS[s["stack"][-1]]
    value = extract(m, s)                           # file_id сохраняется до удаления сообщения
    await try_delete(m)                             # сообщения клиента не остаются в чате
    if value is None:
        return

    if s["ask"]:                                    # клиент отвечает на «Напишите…»
        st = STEPS[s["ask"]["step"]]
        opt = s["ask"]["option"]
        a = s["ans"].setdefault(st["id"], {"title": st["title"], "sel": [], "extra": {}})
        a["extra"][opt["label"]] = value
        s["ask"] = None
        if st["type"] == "multi":
            await show(m, s, st["id"])              # возвращаем то же меню с отметками
        else:
            await go(m, s, next_of(st, opt, s))
    elif step.get("accept_text"):
        if step.get("validate") == "phone" and not 7 <= len(re.sub(r"\D", "", value)) <= 15:
            await show(m, s, step["id"], notice="⚠️ Не похоже на номер телефона. Введите, например: +7 999 123-45-67")
            return
        s["ans"][step["id"]] = {"title": step["title"], "sel": [value], "extra": {}}
        await go(m, s, next_of(step, None, s))


# ───────────────────────── Запуск ─────────────────────────
async def main(dp: Dispatcher):
    load_dotenv(BASE / ".env")
    token = os.getenv("BOT_TOKEN")
    if not token:
        raise RuntimeError("Не задан BOT_TOKEN. Скопируйте .env.example в .env и вставьте токен от @BotFather.")
    bot = Bot(token)
    global runtime_dispatcher
    runtime_dispatcher = dp
    try:
        await bot.set_my_commands([BotCommand(command="start", description="Начать"),
                                   BotCommand(command="contacts", description="Контакты")],
                                  scope=BotCommandScopeDefault())
        await bot.set_chat_menu_button(menu_button=MenuButtonCommands())
        me = await bot.get_me()
        log.info("Бот запущен: @%s. Остановка — Ctrl+C", me.username)
        restart_user_id = settings.get("restart_menu_user")
        if restart_user_id is not None:
            if is_admin(restart_user_id):
                text, markup = admin_home()
                sent = await bot.send_message(restart_user_id, text, reply_markup=markup)
                cleanup = settings.setdefault("chat_cleanup", {})
                ids = cleanup.setdefault(str(restart_user_id), [])
                ids.append(sent.message_id)
                cleanup[str(restart_user_id)] = ids[-100:]
                log.info("Админ-панель открыта после перезапуска для пользователя %s", restart_user_id)
            else:
                log.warning("Не удалось открыть админ-панель: пользователь %s больше не администратор",
                            restart_user_id)
            settings.pop("restart_menu_user", None)
            save_json(SETTINGS_FILE, settings)
        await dp.start_polling(bot, close_bot_session=False)
    finally:
        runtime_dispatcher = None
        await bot.session.close()


async def run():
    global restart_requested
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    dp = Dispatcher()
    dp.include_router(router)
    while True:
        restart_requested = False
        try:
            await main(dp)
            if not restart_requested:
                return
        except Exception:
            log.exception("Бот не запустился или завершился с ошибкой. Повторная попытка через 60 секунд.")
            await asyncio.sleep(60)


if __name__ == "__main__":
    asyncio.run(run())
